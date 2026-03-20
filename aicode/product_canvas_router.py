"""
Product Canvas Router — NPCI Product Build Canvas generator.
Accepts a product document upload, runs deep research via LLM,
and returns a filled 10-section Product Build Canvas.
"""
import json
import os
import re
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from pydantic import BaseModel

# ── LLM ────────────────────────────────────────────────────────────────────
try:
    from langchain_openai import ChatOpenAI
    from langchain_core.messages import SystemMessage, HumanMessage
    _llm = ChatOpenAI(model="gpt-4o-mini", temperature=0.3, max_tokens=6000)
    _LLM_OK = True
except Exception as _e:
    _llm = None
    _LLM_OK = False
    print(f"[Canvas] LLM unavailable: {_e}")

# ── DB ──────────────────────────────────────────────────────────────────────
_DB_PATH = str(Path(__file__).resolve().parent / "artifacts" / "product_canvas.db")
os.makedirs(os.path.dirname(_DB_PATH), exist_ok=True)

router = APIRouter(prefix="/product-canvas", tags=["Product Canvas"])

CANVAS_KEYS = [
    "feature",
    "need_why", "need_differentiation", "need_delta_ux", "need_cannibalize", "need_if_not_built",
    "market_ecosystem_response", "market_ecosystem_efforts", "market_regulatory",
    "scalability_anchors", "scalability_impact",
    "validation_mvp", "validation_data",
    "ops_kpis", "ops_grievance", "ops_day0", "ops_sgf", "ops_frm", "ops_infra",
    "comms_demo", "comms_video", "comms_pm_video", "comms_faqs", "comms_circular", "comms_doc",
    "pricing_3year", "pricing_ability", "pricing_view",
    "risk_fraud", "risk_infosec", "risk_legal", "risk_privacy", "risk_2nd_order",
    "compliance_guideline_change", "compliance_new_guideline", "compliance_npci_circular",
]


def _conn():
    c = sqlite3.connect(_DB_PATH, check_same_thread=False)
    c.row_factory = sqlite3.Row
    return c


def _init_db():
    with _conn() as c:
        c.execute("""
        CREATE TABLE IF NOT EXISTS product_canvas (
            id           TEXT PRIMARY KEY,
            product_name TEXT NOT NULL,
            canvas_data  TEXT NOT NULL DEFAULT '{}',
            status       TEXT NOT NULL DEFAULT 'draft',
            created_at   TEXT NOT NULL,
            updated_at   TEXT NOT NULL
        )
        """)
        c.commit()


_init_db()


# ── Text extraction ─────────────────────────────────────────────────────────
def _extract_text(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="replace")
    if ext == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            return "\n".join(page.get_text() for page in doc)
        except Exception:
            return file_bytes.decode("utf-8", errors="replace")
    if ext == ".docx":
        try:
            from docx import Document
            from io import BytesIO
            doc = Document(BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return file_bytes.decode("utf-8", errors="replace")
    return file_bytes.decode("utf-8", errors="replace")


# ── LLM prompt ──────────────────────────────────────────────────────────────
_SYSTEM_PROMPT = """You are a senior product strategist at NPCI (National Payments Corporation of India) specialising in the UPI payments ecosystem.

Your task: perform DEEP RESEARCH on the product document provided and fill every field of the NPCI Product Build Canvas with thoughtful, specific, and actionable insights.

Rules:
- Reference real payments context: RBI circulars, UPI ecosystem, banks, PSPs, fintechs, SGF, FRM where relevant.
- For each field, provide 2–5 sentences of substantive content — no generic placeholders.
- If specific info is not in the document, apply payments industry best-practice expertise.
- Return ONLY valid JSON with exactly these keys (no markdown fences, no extra keys):

feature, need_why, need_differentiation, need_delta_ux, need_cannibalize, need_if_not_built,
market_ecosystem_response, market_ecosystem_efforts, market_regulatory,
scalability_anchors, scalability_impact,
validation_mvp, validation_data,
ops_kpis, ops_grievance, ops_day0, ops_sgf, ops_frm, ops_infra,
comms_demo, comms_video, comms_pm_video, comms_faqs, comms_circular, comms_doc,
pricing_3year, pricing_ability, pricing_view,
risk_fraud, risk_infosec, risk_legal, risk_privacy, risk_2nd_order,
compliance_guideline_change, compliance_new_guideline, compliance_npci_circular
"""


def _fallback_canvas(product_name: str) -> dict:
    labels = {
        "feature": "Explain the feature for a layman",
        "need_why": "Why should we do this?",
        "need_differentiation": "Differentiation (incremental or exponential)",
        "need_delta_ux": "Delta in user experience",
        "need_cannibalize": "What will it cannibalize?",
        "need_if_not_built": "What if we don't build this?",
        "market_ecosystem_response": "Ecosystem anticipated (informal) response",
        "market_ecosystem_efforts": "Ecosystem efforts (costs to make this work)",
        "market_regulatory": "Anticipated regulatory view",
        "scalability_anchors": "Market anchors to make it big?",
        "scalability_impact": "Impact opportunity (users, delta in time, revenue)",
        "validation_mvp": "Creating and operating MVP",
        "validation_data": "Data it will generate to create insights",
        "ops_kpis": "3 Success KPIs",
        "ops_grievance": "Grievance redressal (Trust)",
        "ops_day0": "Day 0 automation",
        "ops_sgf": "Impact on SGF",
        "ops_frm": "Impact on FRM",
        "ops_infra": "Impact on existing txns and infra",
        "comms_demo": "Product demo (polished version of MVP)",
        "comms_video": "Product video",
        "comms_pm_video": "Explanation video by PM",
        "comms_faqs": "FAQs + trained LLM",
        "comms_circular": "Circular",
        "comms_doc": "Product doc (w/ Specs, explanations, test cases/UI/UX guidelines)",
        "pricing_3year": "3 Year view of pricing & revenue",
        "pricing_ability": "Market ability to pay the price (total pie)",
        "pricing_view": "Market view to pay the price",
        "risk_fraud": "Fraud risk",
        "risk_infosec": "Infosec risk",
        "risk_legal": "Legal risk",
        "risk_privacy": "Data privacy risk",
        "risk_2nd_order": "2nd order negative effect",
        "compliance_guideline_change": "Existing guideline change",
        "compliance_new_guideline": "New guideline addition",
        "compliance_npci_circular": "Must have compliances in NPCI product circular for ecosystem",
    }
    return {k: f"[LLM unavailable — fill in: {v}]" for k, v in labels.items()}


def _generate_canvas(product_name: str, doc_text: str) -> dict:
    user_msg = f"""Product Name: {product_name}

--- PRODUCT DOCUMENT START ---
{doc_text[:10000]}
--- PRODUCT DOCUMENT END ---

Now fill the complete NPCI Product Build Canvas with deep research insights for every section."""

    if _LLM_OK and _llm:
        try:
            result = _llm.invoke([
                SystemMessage(content=_SYSTEM_PROMPT),
                HumanMessage(content=user_msg),
            ])
            raw = result.content.strip()
            raw = re.sub(r"^```[a-z]*\s*", "", raw, flags=re.IGNORECASE)
            raw = re.sub(r"\s*```$", "", raw)
            data = json.loads(raw)
            # Ensure all keys present
            for k in CANVAS_KEYS:
                if k not in data:
                    data[k] = ""
            return data
        except Exception as e:
            print(f"[Canvas LLM error] {e}")
    return _fallback_canvas(product_name)


# ── Endpoints ────────────────────────────────────────────────────────────────

@router.post("/generate")
async def generate_canvas(
    product_name: str = Form(...),
    file: UploadFile = File(...),
):
    file_bytes = await file.read()
    doc_text = _extract_text(file_bytes, file.filename)
    canvas_data = _generate_canvas(product_name, doc_text)

    canvas_id = str(uuid.uuid4())
    now = datetime.utcnow().isoformat()
    with _conn() as c:
        c.execute(
            "INSERT INTO product_canvas(id, product_name, canvas_data, status, created_at, updated_at) "
            "VALUES(?,?,?,?,?,?)",
            (canvas_id, product_name, json.dumps(canvas_data), "draft", now, now),
        )
        c.commit()

    return {
        "id": canvas_id,
        "product_name": product_name,
        "canvas": canvas_data,
        "status": "draft",
        "created_at": now,
    }


@router.get("/list")
def list_canvases():
    with _conn() as c:
        rows = c.execute(
            "SELECT id, product_name, status, created_at, updated_at FROM product_canvas ORDER BY created_at DESC"
        ).fetchall()
    return [dict(r) for r in rows]


@router.get("/{canvas_id}")
def get_canvas(canvas_id: str):
    with _conn() as c:
        row = c.execute("SELECT * FROM product_canvas WHERE id=?", (canvas_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Canvas not found")
    d = dict(row)
    d["canvas"] = json.loads(d["canvas_data"])
    return d


@router.patch("/{canvas_id}")
async def update_canvas(canvas_id: str, body: dict):
    with _conn() as c:
        row = c.execute("SELECT canvas_data, product_name FROM product_canvas WHERE id=?", (canvas_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Canvas not found")
        canvas_data = json.loads(row["canvas_data"])
        product_name = row["product_name"]
        if "canvas" in body:
            canvas_data.update(body["canvas"])
        if "product_name" in body:
            product_name = body["product_name"]
        now = datetime.utcnow().isoformat()
        c.execute(
            "UPDATE product_canvas SET product_name=?, canvas_data=?, updated_at=? WHERE id=?",
            (product_name, json.dumps(canvas_data), now, canvas_id),
        )
        c.commit()
    return {"success": True}


@router.patch("/{canvas_id}/status")
async def update_canvas_status(canvas_id: str, body: dict):
    with _conn() as c:
        row = c.execute("SELECT id FROM product_canvas WHERE id=?", (canvas_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Canvas not found")
        now = datetime.utcnow().isoformat()
        c.execute(
            "UPDATE product_canvas SET status=?, updated_at=? WHERE id=?",
            (body.get("status", "draft"), now, canvas_id),
        )
        c.commit()
    return {"success": True}


@router.delete("/{canvas_id}")
def delete_canvas(canvas_id: str):
    with _conn() as c:
        c.execute("DELETE FROM product_canvas WHERE id=?", (canvas_id,))
        c.commit()
    return {"success": True}
